"""Generate a single-page maths exercise sheet (.docx) from a CSV file.

CSV columns: Number,ExerciseName,LaTeX1,...,LaTeXN (or a single legacy LaTeX column).
Each nonempty LaTeX field becomes a left-aligned equation line; '--blank--' creates a blank line.
"""

import argparse
import csv
import os
import sys
from pathlib import Path
from typing import Any, cast

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from latex2mathml.converter import convert as latex_to_mathml
from lxml import etree  # type: ignore[import-not-found]

FONT_NAME = "Arial"
TITLE_SIZE = 14
MAX_FONT_SIZE = 12
MIN_FONT_SIZE = 8
MIN_GAP_PT = 4
MAX_GAP_PT = 20

PAGE_WIDTH = Cm(21.0)
PAGE_HEIGHT = Cm(29.7)
MARGIN = Cm(1.5)

XSL_CANDIDATES = [
    r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL",
    r"C:\Program Files (x86)\Microsoft Office\root\Office16\MML2OMML.XSL",
    r"C:\Program Files\Microsoft Office\Office16\MML2OMML.XSL",
    r"C:\Program Files (x86)\Microsoft Office\Office16\MML2OMML.XSL",
    r"C:\Program Files\Microsoft Office\Office15\MML2OMML.XSL",
]


def find_xsl(explicit):
    candidates = [explicit] if explicit else [os.environ.get("MML2OMML_XSL"), *XSL_CANDIDATES]
    for path in candidates:
        if path and Path(path).is_file():
            return path
    sys.exit("MML2OMML.XSL not found. Pass --xsl <path> or set MML2OMML_XSL.")


def read_exercises(csv_path):
    with open(csv_path, newline="", encoding="utf-8-sig") as f:
        rows = []
        for row in csv.DictReader(f):
            row = {k.strip(): (v or "").strip() for k, v in row.items() if k}
            if not exercise_lines(row):
                continue
            rows.append(row)
    return rows


def exercise_lines(exercise):
    lines = []
    for key, value in exercise.items():
        suffix = "1" if key == "LaTeX" else key.removeprefix("LaTeX")
        if key.startswith("LaTeX") and suffix.isdecimal() and int(suffix) >= 1:
            latex = value.strip()
            if latex:
                lines.append((int(suffix), latex))
    return sorted(lines)


def clean_latex(latex):
    latex = latex.strip()
    for delim in ("$$", "$"):
        if latex.startswith(delim) and latex.endswith(delim) and len(latex) > 2 * len(delim):
            return latex[len(delim):-len(delim)].strip()
    if latex.startswith(r"\(") and latex.endswith(r"\)"):
        return latex[2:-2].strip()
    return latex


def latex_to_omml(latex, transform):
    mathml = etree.fromstring(latex_to_mathml(clean_latex(latex)).encode("utf-8"))
    omml = transform(mathml).getroot()
    if omml.tag == qn("m:oMath"):
        math_paragraph = OxmlElement("m:oMathPara")
        math_paragraph.append(omml)
        omml = math_paragraph
    properties = omml.find(qn("m:oMathParaPr"))
    if properties is None:
        properties = OxmlElement("m:oMathParaPr")
        omml.insert(0, properties)
    justification = properties.find(qn("m:jc"))
    if justification is None:
        justification = OxmlElement("m:jc")
        properties.append(justification)
    justification.set(qn("m:val"), "left")
    return omml


def estimate_exercise_height(exercise, font_pt):
    line = font_pt * 1.2
    lines = 0
    for _, latex in exercise_lines(exercise):
        if latex == "--blank--":
            lines += 1
        else:
            lines += 2.2 if any(t in latex for t in (r"\frac", r"\dfrac", r"\sqrt", r"\sum", r"\int", "^", "_")) else 1.2
    return line * lines


def choose_layout(exercises, available_pt):
    """Pick the largest font size and paragraph gap that fit on one page."""
    for font_pt in range(MAX_FONT_SIZE, MIN_FONT_SIZE - 1, -1):
        worst = sum(estimate_exercise_height(exercise, font_pt) for exercise in exercises)
        rows = sum(len(exercise_lines(exercise)) for exercise in exercises)
        gap = (available_pt - worst) / max(rows, 1)
        if gap >= MIN_GAP_PT:
            return font_pt, min(gap, MAX_GAP_PT)
    sys.exit("Too many exercises to fit on a single page.")


def build_document(title, exercises, transform, output):
    doc = Document()
    normal = cast(Any, doc.styles["Normal"])
    normal.font.name = FONT_NAME
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)

    section = doc.sections[0]
    section.page_width, section.page_height = PAGE_WIDTH, PAGE_HEIGHT
    section.top_margin = MARGIN
    section.bottom_margin = MARGIN - Pt(MAX_FONT_SIZE)
    section.left_margin = section.right_margin = MARGIN

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(12)
    run = title_p.add_run(title)
    run.font.name = FONT_NAME
    run.font.size = Pt(TITLE_SIZE)
    run.bold = True

    available_pt = PAGE_HEIGHT.pt - 2 * MARGIN.pt - (TITLE_SIZE * 1.2 + 12) - 6
    font_pt, gap = choose_layout(exercises, available_pt)
    normal.font.size = Pt(font_pt)

    for exercise in exercises:
        lines = exercise_lines(exercise)
        for line_index, (_, latex) in enumerate(lines):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(gap)
            paragraph.paragraph_format.keep_with_next = line_index < len(lines) - 1
            if latex != "--blank--":
                paragraph._p.append(latex_to_omml(latex, transform))

    doc.save(output)


def main():
    parser = argparse.ArgumentParser(description="Create a single-page maths exercise sheet (.docx) from a CSV file.")
    parser.add_argument("title", help="Sheet title")
    parser.add_argument("csv_file", help="CSV with columns: Number,ExerciseName,LaTeX1,...,LaTeXN or LaTeX (use --blank-- for a blank line)")
    parser.add_argument("-o", "--output", help="Output .docx path (default: CSV name with .docx extension)")
    parser.add_argument("--xsl", help="Path to MML2OMML.XSL (from Microsoft Office)")
    args = parser.parse_args()

    exercises = read_exercises(args.csv_file)
    if not exercises:
        sys.exit("No exercises found in CSV.")

    transform = etree.XSLT(etree.parse(find_xsl(args.xsl)))
    output = args.output or str(Path(args.csv_file).with_suffix(".docx"))
    build_document(args.title, exercises, transform, output)
    print(f"Created {output}")


if __name__ == "__main__":
    main()
